#!/usr/bin/env python3
"""
Excel全シートをマークダウンテーブルに変換するスクリプト
全シートを変換してシートごとにmdファイルを出力する
出力先: {outdir}/{Excelファイル名}/{n}-{シート名}.md
"""
import argparse
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("エラー: openpyxlがインストールされていません。")
    print("以下のコマンドでインストールしてください:")
    print("  pip3 install openpyxl")
    sys.exit(1)


def cell_to_str(cell_value) -> str:
    if cell_value is None:
        return ""
    if isinstance(cell_value, float) and cell_value == int(cell_value):
        s = str(int(cell_value))
    else:
        s = str(cell_value)
    s = s.replace("\n", "<br>")
    s = re.sub(r'・', '- ', s)
    s = re.sub(r'(\d+)\.', r'\1. ', s)
    return s


def sheet_to_markdown(ws, sheet_name: str, ref_id: str, ref_title: str) -> str:
    """
    1シートをマークダウンテーブルに変換する

    - 全行・全列を動的に読み込む
    - 最初の非空行をヘッダーとして扱う
    - 2列目・3列目の〃（同上マーク）は直前行の値に復元してから再度〃表示する
    """
    all_rows: list[list[str]] = []
    for row in ws.iter_rows():
        row_data = [cell_to_str(c.value) for c in row]
        while row_data and row_data[-1] == "":
            row_data.pop()
        if any(cell != "" for cell in row_data):
            all_rows.append(row_data)

    if not all_rows:
        return f"# {sheet_name}\n\n（データなし）\n"

    max_col = max(len(r) for r in all_rows)
    for r in all_rows:
        while len(r) < max_col:
            r.append("")

    if ref_id and ref_title:
        title = f"# {sheet_name}（#{ref_id} - {ref_title}）"
    else:
        title = f"# {sheet_name}"

    lines = [title, ""]
    header = all_rows[0]
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join(["---"] * max_col) + " |")

    prev_row: list[str] | None = None
    for row in all_rows[1:]:
        display = list(row)
        for col_idx in [1, 2]:  # 2列目・3列目
            if col_idx < max_col:
                actual_value = row[col_idx]
                if actual_value == "〃" and prev_row and col_idx < len(prev_row):
                    actual_value = prev_row[col_idx]
                    row[col_idx] = actual_value

                if prev_row and actual_value and actual_value == prev_row[col_idx]:
                    display[col_idx] = "〃"
                else:
                    display[col_idx] = actual_value

        lines.append("| " + " | ".join(display) + " |")
        prev_row = row

    return "\n".join(lines) + "\n"


def safe_filename(name: str) -> str:
    return re.sub(r'[/\\:*?"<>|]', '_', name)


def main():
    parser = argparse.ArgumentParser(description='Excel全シートをマークダウンに変換')
    parser.add_argument('--input', required=True, help='入力Excelファイルパス')
    parser.add_argument('--outdir', default='temp', help='親出力ディレクトリ（デフォルト: temp）')
    parser.add_argument('--ref-id', default='', help='見出しに付与する参照ID（チケット番号など、省略可）')
    parser.add_argument('--ref-title', default='', help='見出しに付与する参照タイトル（省略可）')
    parser.add_argument('--sheet', default='', help='特定シートのみ変換（省略時は全シート）')

    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"エラー: 入力ファイルが見つかりません: {args.input}", file=sys.stderr)
        sys.exit(1)

    # Excelファイル名（拡張子なし）をサブディレクトリ名として使用
    excel_stem = input_path.stem
    outdir = Path(args.outdir) / safe_filename(excel_stem)
    outdir.mkdir(parents=True, exist_ok=True)

    print(f"📊 Excelファイルを読み込み中: {args.input}")
    wb = openpyxl.load_workbook(args.input, data_only=True)

    # 表示されているシートのみ対象（hidden/veryHidden は除外）
    visible_sheets = [
        name for name in wb.sheetnames
        if wb[name].sheet_state == 'visible'
    ]
    target_sheets = [args.sheet] if args.sheet else visible_sheets

    total_rows = 0
    for idx, sheet_name in enumerate(target_sheets, start=1):
        if sheet_name not in wb.sheetnames:
            print(f"  ⚠️  シートが見つかりません: {sheet_name}", file=sys.stderr)
            continue

        ws = wb[sheet_name]
        md = sheet_to_markdown(ws, sheet_name, args.ref_id, args.ref_title)

        # {n}-{シート名}.md の形式でファイル名を生成
        filename = f"{idx}-{safe_filename(sheet_name)}.md"
        out_path = outdir / filename
        out_path.write_text(md, encoding='utf-8')

        data_rows = md.count('\n|') - 2
        total_rows += max(data_rows, 0)
        print(f"  ✅ [{idx}] {sheet_name} → {out_path}（{max(data_rows, 0)}行）")

    print(f"\n✅ 変換完了: {len(target_sheets)}シート / 合計{total_rows}行")
    print(f"   出力先: {outdir}")


if __name__ == '__main__':
    main()
